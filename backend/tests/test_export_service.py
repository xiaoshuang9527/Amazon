"""导出服务测试：xlsx（Flat File）/ json（SP-API patch）/ csv。"""

import csv
import json

import pytest
from openpyxl import load_workbook

from app.core.exceptions import AppError
from app.services.export_service import FLAT_FILE_HEADERS, export, export_csv, export_json, export_xlsx

TASK_ID = "taskexport01"

RESULTS = [
    {
        "target_marketplace": "US",
        "target_language": "en",
        "title": "Wireless Bluetooth Headphones with 40H Playtime",
        "bullets": ["b1", "b2", "b3", "b4", "b5"],
        "description": "English description",
        "search_terms": "wireless headphones bluetooth",
        "suggested_keywords": ["wireless headphones"],
        "compliance_report": {"passed": True, "violations": []},
    },
    {
        "target_marketplace": "DE",
        "target_language": "de",
        "title": "Kabellose Bluetooth Kopfhörer",
        "bullets": ["d1", "d2"],
        "description": "Deutsche Beschreibung",
        "search_terms": "kopfhörer bluetooth",
        "suggested_keywords": ["kopfhörer"],
        "compliance_report": {
            "passed": False,
            "violations": [{"type": "absolute_term", "word": "beste", "message": "命中绝对化用语"}],
        },
    },
]

ROWS = [
    {"target_marketplace": "US", "sku": "SKU-US-1"},
    {"target_marketplace": "DE", "sku": "SKU-DE-1"},
]


# --------------------------------------------------------------------------- #
# JSON
# --------------------------------------------------------------------------- #
def test_export_json_follows_spapi_patch_shape(tmp_path):
    path = export_json(TASK_ID, RESULTS, ROWS, str(tmp_path))
    payload = json.loads(open(path, encoding="utf-8").read())

    assert payload["task_id"] == TASK_ID
    assert len(payload["items"]) == 2

    us_item = payload["items"][0]
    assert us_item["sku"] == "SKU-US-1"
    assert us_item["marketplace_id"] == "US"
    assert us_item["operation_type"] == "PATCH"
    assert us_item["patches"][0]["op"] == "replace"
    assert us_item["patches"][0]["path"] == "/attributes"

    attributes = us_item["patches"][0]["value"]
    assert attributes["item_name"][0]["language_tag"] == "en"
    assert attributes["item_name"][0]["value"] == RESULTS[0]["title"]
    assert len(attributes["bullet_point"]) == 5
    assert attributes["generic_keyword"][0]["value"] == RESULTS[0]["search_terms"]


def test_export_json_uses_placeholder_sku_when_missing(tmp_path):
    path = export_json(TASK_ID, RESULTS, [], str(tmp_path))
    payload = json.loads(open(path, encoding="utf-8").read())
    assert payload["items"][0]["sku"] == f"TASK-{TASK_ID}"


def test_export_json_keeps_compliance_report(tmp_path):
    path = export_json(TASK_ID, RESULTS, ROWS, str(tmp_path))
    payload = json.loads(open(path, encoding="utf-8").read())
    assert payload["items"][1]["compliance_report"]["passed"] is False


# --------------------------------------------------------------------------- #
# CSV
# --------------------------------------------------------------------------- #
def test_export_csv_has_flat_file_headers_and_rows(tmp_path):
    path = export_csv(TASK_ID, RESULTS, ROWS, str(tmp_path))
    with open(path, encoding="utf-8-sig", newline="") as fp:
        rows = list(csv.reader(fp))

    assert rows[0] == FLAT_FILE_HEADERS
    assert len(rows) == 3

    us_row = rows[1]
    assert us_row[0] == "SKU-US-1"
    assert us_row[1] == RESULTS[0]["title"]
    assert us_row[3:8] == ["b1", "b2", "b3", "b4", "b5"]
    assert us_row[9] == "US" and us_row[10] == "en"


def test_export_csv_pads_missing_bullets(tmp_path):
    path = export_csv(TASK_ID, RESULTS, ROWS, str(tmp_path))
    with open(path, encoding="utf-8-sig", newline="") as fp:
        rows = list(csv.reader(fp))
    de_row = rows[2]
    assert de_row[3:8] == ["d1", "d2", "", "", ""]


# --------------------------------------------------------------------------- #
# XLSX
# --------------------------------------------------------------------------- #
def test_export_xlsx_creates_listing_and_compliance_sheets(tmp_path):
    path = export_xlsx(TASK_ID, RESULTS, ROWS, str(tmp_path))
    wb = load_workbook(path)

    assert wb.sheetnames == ["Listing", "Compliance"]

    ws = wb["Listing"]
    header = [cell.value for cell in ws[1]]
    assert header == FLAT_FILE_HEADERS
    assert ws.max_row == 3
    assert ws.cell(row=2, column=1).value == "SKU-US-1"
    assert ws.cell(row=2, column=10).value == "US"
    assert ws.cell(row=3, column=10).value == "DE"

    ws2 = wb["Compliance"]
    assert [c.value for c in ws2[1]] == ["marketplace", "passed", "violations"]
    assert ws2.cell(row=2, column=2).value is True
    assert ws2.cell(row=3, column=2).value is False
    assert "绝对化用语" in ws2.cell(row=3, column=3).value


# --------------------------------------------------------------------------- #
# 统一入口
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    "fmt,expected_suffix",
    [("xlsx", ".xlsx"), ("json", ".json"), ("csv", ".csv"), ("XLSX", ".xlsx")],
)
def test_export_dispatch(tmp_path, fmt, expected_suffix):
    path = export(fmt, TASK_ID, RESULTS, ROWS, str(tmp_path))
    assert path.endswith(f"listing_{TASK_ID}{expected_suffix}")


def test_export_rejects_unknown_format(tmp_path):
    with pytest.raises(AppError):
        export("pdf", TASK_ID, RESULTS, ROWS, str(tmp_path))


def test_export_creates_missing_directory(tmp_path):
    nested = tmp_path / "a" / "b"
    path = export("csv", TASK_ID, RESULTS, ROWS, str(nested))
    assert nested.exists() and path.endswith(".csv")
